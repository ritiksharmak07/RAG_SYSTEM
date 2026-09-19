from pathlib import Path

from docx import Document as DocxDocument

from utils.document import Document


class DOCXLoader:

    def load(self, file_path: str):

        file_path = Path(file_path)

        doc = DocxDocument(file_path)

        text = "\n".join([p.text for p in doc.paragraphs])

        return [

            Document(

                text=text,

                metadata={
                    "filename": file_path.name,
                    "filepath": str(file_path),
                    "source": "docx"
                }

            )

        ]